import pandas as pd
import numpy as np
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

np.random.seed(42)

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
data_font = Font(name="Segoe UI", size=10)
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# -------------------------------
# 1. MES EQUIPMENT OPERATION RATE
# -------------------------------
dates = pd.date_range(start="2026-08-01", end="2026-08-31", freq="D")
equipment_list = [
    "auto label #1", "auto trimming #1", "backend el #1",
    "curing room #1", "frame assembly #1", "laminator #1",
    "tabber #1", "tabber #2", "matrix layup #1", "string repair #1"
]

eqp_records = []
for dt in dates:
    day_perf = np.random.choice(["Under Expectations", "Meeting Expectations", "Above Target"], p=[0.25, 0.50, 0.25])
    for eqp in equipment_list:
        if day_perf == "Under Expectations":
            run_pct = np.random.uniform(40.0, 58.0)
            down_machine = np.random.uniform(15.0, 30.0)
            no_wip = np.random.uniform(15.0, 25.0)
        elif day_perf == "Meeting Expectations":
            run_pct = np.random.uniform(65.0, 78.0)
            down_machine = np.random.uniform(2.0, 10.0)
            no_wip = np.random.uniform(10.0, 20.0)
        else:
            run_pct = np.random.uniform(82.0, 95.0)
            down_machine = np.random.uniform(0.5, 5.0)
            no_wip = np.random.uniform(2.0, 10.0)

        remaining = max(0.0, 100.0 - (run_pct + down_machine + no_wip))
        kanban_full = remaining * 0.5
        tool_pause = remaining * 0.3
        other_losses = remaining * 0.2

        eqp_records.append({
            "Work Date": dt.strftime("%Y-%m-%d"),
            "Equipment": eqp,
            "Performance Status": day_perf,
            "RUN Rate (%)": round(run_pct / 100.0, 4),
            "No WIP (%)": round(no_wip / 100.0, 4),
            "Down Machine (%)": round(down_machine / 100.0, 4),
            "Kanban Full (%)": round(kanban_full / 100.0, 4),
            "Tool Pause (%)": round(tool_pause / 100.0, 4),
            "Other Losses (%)": round(other_losses / 100.0, 4)
        })

df_eqp = pd.DataFrame(eqp_records)
file_eqp = "MES_Equipment_Operation_Rate.xlsx"

with pd.ExcelWriter(file_eqp, engine='openpyxl') as writer:
    df_eqp.to_excel(writer, sheet_name="Equipment Operation Rates", index=False)

wb_eqp = openpyxl.load_workbook(file_eqp)
ws_eqp = wb_eqp["Equipment Operation Rates"]

for col_num in range(1, len(df_eqp.columns) + 1):
    cell = ws_eqp.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row_idx in range(2, len(df_eqp) + 2):
    for col_num in range(1, len(df_eqp.columns) + 1):
        cell = ws_eqp.cell(row=row_idx, column=col_num)
        cell.font = data_font
        cell.border = thin_border
        if col_num in [1, 2, 3]:
            cell.alignment = Alignment(horizontal="center")
        else:
            cell.number_format = '0.00%'
            cell.alignment = Alignment(horizontal="right")

for col in ws_eqp.columns:
    max_len = max(len(str(cell.value or '')) for cell in col[:100])
    col_letter = get_column_letter(col[0].column)
    ws_eqp.column_dimensions[col_letter].width = max(max_len + 3, 14)

ws_eqp.freeze_panes = "A2"
wb_eqp.save(file_eqp)

# ----------------------------
# 2. SAP PRODUCTION VOLUME
# ----------------------------
line_models = [
    {"Line": "Line A", "Model": "M-100", "Target": 42000, "Factor": 1.00},  # Meeting target (100%)
    {"Line": "Line B", "Model": "M-200", "Target": 42000, "Factor": 0.80},  # Underperforming (80%)
    {"Line": "Line C", "Model": "M-300", "Target": 42000, "Factor": 1.057}  # Above target (105.7%)
]

power_classes = [585, 590, 595, 650, 655, 660, 675]
weekly_dates = pd.date_range(start="2026-08-01", end="2026-08-07", freq="D")
prod_records = []

for dt in weekly_dates:
    day_weight = 1.0 + np.random.uniform(-0.03, 0.03)
    for lm in line_models:
        p_class = np.random.choice(power_classes)
        base_daily_pcs = (lm["Target"] / 7) * lm["Factor"] * day_weight
        pcs_produced = int(round(base_daily_pcs))
        mw_produced = round((pcs_produced * p_class) / 1000000.0, 4)

        prod_records.append({
            "Posting Date": dt.strftime("%Y-%m-%d"),
            "Line": lm["Line"],
            "Material Code": lm["Model"],
            "Power Class (W)": p_class,
            "Production Volume (PCS)": pcs_produced,
            "Capacity (MW)": mw_produced
        })

df_prod = pd.DataFrame(prod_records)
file_prod = "SAP_Production_Volume.xlsx"

with pd.ExcelWriter(file_prod, engine='openpyxl') as writer:
    df_prod.to_excel(writer, sheet_name="Production Volume", index=False)

wb_prod = openpyxl.load_workbook(file_prod)
ws_prod = wb_prod["Production Volume"]

for col_num in range(1, len(df_prod.columns) + 1):
    cell = ws_prod.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row_idx in range(2, len(df_prod) + 2):
    for col_num in range(1, len(df_prod.columns) + 1):
        cell = ws_prod.cell(row=row_idx, column=col_num)
        cell.font = data_font
        cell.border = thin_border
        if col_num in [1, 2, 3, 4]:
            cell.alignment = Alignment(horizontal="center")
        elif col_num == 5:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right")
        elif col_num == 6:
            cell.number_format = '#,##0.0000'
            cell.alignment = Alignment(horizontal="right")

for col in ws_prod.columns:
    max_len = max(len(str(cell.value or '')) for cell in col[:100])
    col_letter = get_column_letter(col[0].column)
    ws_prod.column_dimensions[col_letter].width = max(max_len + 3, 14)

ws_prod.freeze_panes = "A2"
wb_prod.save(file_prod)

# -------------------------------------
# 3. SAP MATERIAL CONSUMPTION
# -------------------------------------
num_days = 7
records_per_day = 3448
total_sap_rows = num_days * records_per_day

plants = ["P313"]
storage_locs = ["LRH2"]
mvt_types = ["311"]
mvt_texts = ["TF all within plant"]
materials = [
    ("11800110", "QHQUP15-G2-01 135/190 US"),
    ("11800120", "R G/CSG/ARC/045T/1126/2.0TM16/G/CH"),
    ("CF020G02-0103", "Silicone/Potting/CV115-A/WH/5KG/A/CH"),
    ("CF020G02-0104", "Silicone/Potting/CV115-B/WH/5KG/B/CH"),
    ("12200030", "Frame Corner Cap Assembly")
]
units = ["PCS", "KG", "EA"]

mat_choices = [materials[i] for i in np.random.choice(len(materials), total_sap_rows)]
mat_codes_cons = [m[0] for m in mat_choices]
mat_descs_cons = [m[1] for m in mat_choices]

posting_dates = [(datetime(2026, 8, 1) + timedelta(days=i // records_per_day)).strftime("%Y-%m-%d") for i in
                 range(total_sap_rows)]
quantities = np.random.randint(10, 500, total_sap_rows)
selected_units = np.random.choice(units, total_sap_rows)

df_cons = pd.DataFrame({
    "Plant": np.random.choice(plants, total_sap_rows),
    "Stor. Loc": np.random.choice(storage_locs, total_sap_rows),
    "Movement Type": np.random.choice(mvt_types, total_sap_rows),
    "Movement Type Text": np.random.choice(mvt_texts, total_sap_rows),
    "Material Code": mat_codes_cons,
    "Material Description": mat_descs_cons,
    "Quantity": quantities,
    "Base Unit of Measure": selected_units,
    "Posting Date": posting_dates,
    "Amount Local Curr": 0.00,
    "Currency": "USD"
})

file_cons = "SAP_Material_Consumption.xlsx"

with pd.ExcelWriter(file_cons, engine='openpyxl') as writer:
    df_cons.to_excel(writer, sheet_name="Material Movement", index=False)

wb_cons = openpyxl.load_workbook(file_cons)
ws_cons = wb_cons["Material Movement"]

for col_num in range(1, len(df_cons.columns) + 1):
    cell = ws_cons.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for col in ws_cons.columns:
    max_len = max(len(str(cell.value or '')) for cell in col[:100])
    col_letter = get_column_letter(col[0].column)
    ws_cons.column_dimensions[col_letter].width = max(max_len + 3, 14)

ws_cons.freeze_panes = "A2"
wb_cons.save(file_cons)